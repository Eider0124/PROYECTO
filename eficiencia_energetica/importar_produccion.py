"""Extrae la producción de la línea 22 del Excel 'KG producidos 2024, 2025 y 2026'
(hoja 'Data PA 2026', un registro por fecha + turno + SKU) y la guarda en datos/.

Uso:  python3 importar_produccion.py <ruta del Excel>

Genera:
  datos/produccion_l22_detalle.csv   un registro por fecha, turno y SKU
  datos/produccion_l22_diaria.csv    kg por día y por turno
Después se corre construir_base.py para cargarla en mediciones y en la base SQLite.
"""
import csv, sys, collections
from pathlib import Path
from openpyxl import load_workbook

BASE = Path(__file__).parent
LINEA = "LINEA 22"


def num(x):
    return x if isinstance(x, (int, float)) else 0.0


def main(ruta):
    wb = load_workbook(ruta, read_only=True, data_only=True)
    ws = wb["Data PA 2026"]
    filas = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and r[1] == LINEA]
    datos = BASE / "datos"
    with open(datos / "produccion_l22_detalle.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["fecha", "turno", "sku", "descripcion", "cajas", "embalaje", "unidades", "kg_caja", "kg", "mantenimiento"])
        for f, _, turno, sku, desc, cajas, _, emb, und, kgc, kg in sorted(filas, key=lambda r: (r[0], r[2], r[3])):
            w.writerow([f.date().isoformat(), turno, sku, (desc or "").strip(), num(cajas),
                        emb if isinstance(emb, (int, float)) else "", num(und),
                        kgc if isinstance(kgc, (int, float)) else "", round(num(kg), 3),
                        "si" if sku == 22 else "no"])
    dia = collections.defaultdict(lambda: {"t1": 0.0, "t2": 0.0, "mant": False})
    for r in filas:
        d = dia[r[0].date()]
        d["t1" if r[2] == 1 else "t2"] += num(r[10])
        d["mant"] |= r[3] == 22
    with open(datos / "produccion_l22_diaria.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["fecha", "kg_turno1", "kg_turno2", "kg_total", "mantenimiento"])
        for f in sorted(dia):
            d = dia[f]
            w.writerow([f.isoformat(), round(d["t1"], 3), round(d["t2"], 3),
                        round(d["t1"] + d["t2"], 3), "si" if d["mant"] else "no"])
    print(f"{len(filas)} registros, {len(dia)} días de producción de la L22")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
